# -*- coding: utf-8 -*-
"""
Генератор демо-таблицы для кейса «Excel с продажами»
(МК «Автоматизация бизнес-процессов с AI»).

Создаёт рядом с собой:
  «Продажи лето 2026 ФИНАЛ (2).xlsx» — грязная таблица трёх менеджеров, её показываем на МК;
  «_эталон.xlsx»                      — правильный ответ: чистые заявки, воронки, скорость КП,
                                        ловушки и сверка. По нему проверяем результат AI.

Всё детерминировано (SEED): при каждом запуске получается один и тот же файл.
Запуск: python generate_sales.py
"""
import datetime as dt
import random
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
DIRTY_FILE = HERE / "Продажи лето 2026 ФИНАЛ (2).xlsx"
TRUTH_FILE = HERE / "_эталон.xlsx"

SEED = 2026
EUR_RATE = 20.02                     # курс для демо, 1 EUR = 20,02 MDL
SNAPSHOT = dt.date(2026, 9, 7)       # «сегодня» в легенде: понедельник, бухгалтер сводит отчёт
PERIOD_START = dt.date(2026, 6, 1)
PERIOD_END = dt.date(2026, 8, 23)
MONTHS_IN_PERIOD = 3

rng = random.Random(SEED)

# ---------------------------------------------------------------------------
# Параметры бизнеса: здесь заложены инсайты, которые должна найти воронка
# ---------------------------------------------------------------------------

CHANNEL_COUNTS = {"Instagram": 158, "999.md": 94, "Рекомендация": 52,
                  "Звонок": 48, "Facebook": 44, "Не указан": 14}
P_MEASURE = {"Instagram": 0.36, "999.md": 0.55, "Рекомендация": 0.85,
             "Звонок": 0.65, "Facebook": 0.42, "Не указан": 0.50}
P_KP_AFTER_MEASURE = 0.88
BASE_CONTRACT = {"Instagram": 0.30, "999.md": 0.38, "Рекомендация": 0.52,
                 "Звонок": 0.42, "Facebook": 0.30, "Не указан": 0.38}
SPEED_FACTOR = {True: 1.45, False: 0.60}   # КП за сутки после замера против медленного
P_PAY_AFTER_CONTRACT = 0.96

MANAGERS = {"Наталья": 0.36, "Ion": 0.34, "Андрей": 0.30}
P_FAST_KP = {"Наталья": 0.80, "Ion": 0.25, "Андрей": 0.50}
DATIVE = {"Наталья": "Наталье", "Ion": "Иону", "Андрей": "Андрею"}

STAGES = ["Заявка", "Замер", "КП", "Договор", "Оплата"]

# ---------------------------------------------------------------------------
# Справочники: канон → варианты написания у каждого менеджера
# ---------------------------------------------------------------------------

CITIES = {
    "Кишинёв": (55, {"Наталья": ["Кишинёв", "Кишинев", "Кишинёв, Ботаника", "Кишинев (Рышкановка)"],
                     "Ion": ["Chișinău", "Chisinau", "Chișinău, sec. Centru", "Кишинев"],
                     "Андрей": ["кшн", "КИШИНЕВ", "Кишинев ", "Kишинёв", "Кишинёв", "Ботаника"]}),
    "Бельцы": (9, {"Наталья": ["Бельцы"], "Ion": ["Bălți", "Balti"], "Андрей": ["Бельцы", "Бэлць"]}),
    "Дурлешты": (5, {"Наталья": ["Дурлешты"], "Ion": ["Durlești", "Durlesti"], "Андрей": ["Дурлешты"]}),
    "Яловены": (5, {"Наталья": ["Яловены"], "Ion": ["Ialoveni"], "Андрей": ["Яловены", "яловены"]}),
    "Орхей": (5, {"Наталья": ["Орхей"], "Ion": ["Orhei"], "Андрей": ["Оргеев", "Орхей"]}),
    "Страшены": (4, {"Наталья": ["Страшены"], "Ion": ["Strășeni", "Straseni"], "Андрей": ["Страшены"]}),
    "Унгены": (4, {"Наталья": ["Унгены"], "Ion": ["Ungheni"], "Андрей": ["Унгены"]}),
    "Криково": (4, {"Наталья": ["Криково"], "Ion": ["Cricova"], "Андрей": ["Криково"]}),
    "Кэушень": (3, {"Наталья": ["Кэушень"], "Ion": ["Căușeni"], "Андрей": ["Каушаны"]}),
    "Комрат": (3, {"Наталья": ["Комрат"], "Ion": ["Comrat"], "Андрей": ["Комрат"]}),
    "Ставчены": (3, {"Наталья": ["Ставчены"], "Ion": ["Stăuceni"], "Андрей": ["Ставчены"]}),
}

SOURCES = {
    "Instagram": {"Наталья": ["Instagram", "Instagram", "инста"], "Ion": ["Insta", "IG", "Instagram"],
                  "Андрей": ["инстаграм", "IG", "insta", "Instagram "]},
    "999.md": {"Наталья": ["999.md"], "Ion": ["999", "999.md"], "Андрей": ["с 999", "999", "объявление 999"]},
    "Рекомендация": {"Наталья": ["Рекомендация", "По рекомендации"], "Ion": ["recomandare", "рекомендация"],
                     "Андрей": ["от Ивана", "знакомый", "сосед посоветовал", "рек."]},
    "Звонок": {"Наталья": ["Звонок"], "Ion": ["apel", "звонок"], "Андрей": ["позвонил", "тел.", "звонила"]},
    "Facebook": {"Наталья": ["Facebook"], "Ion": ["FB", "Facebook"], "Андрей": ["фейсбук", "fb", "Facebook реклама"]},
    "Не указан": {"Наталья": [""], "Ion": [""], "Андрей": [""]},
}

STATUSES = {
    "Наталья": {"Оплачено": ["Оплачено"], "Договор": ["Договор"], "КП отправлено": ["КП отправлено"],
                "Думает": ["Думает"], "Отказ": ["Отказ"], "Не дозвонились": ["Не дозвонилась"]},
    "Ion": {"Оплачено": ["achitat", "оплачено", "achitat ✓"], "Договор": ["contract", "договор"],
            "КП отправлено": ["ofertă trimisă"], "Думает": ["se gândește", "думает"],
            "Отказ": ["refuz", "отказ"], "Не дозвонились": ["nu răspunde"]},
    "Андрей": {"Оплачено": ["+", "да", "опл", "оплатили"], "Договор": ["дог.", "договор подписан"],
               "КП отправлено": ["кп"], "Думает": ["думает", "перезвонить", "думают"],
               "Отказ": ["нет", "пропал", "отказ"], "Не дозвонились": ["не дозвон", "пропал"]},
}

# канон: (вес, диапазон суммы в MDL, варианты по менеджерам)
PRODUCTS = {
    "Окна в квартиру": (30, (14000, 42000), {
        "Наталья": ["Окна в квартиру", "Окна, квартира (3 шт)"],
        "Ion": ["ferestre apartament", "окна квартира"],
        "Андрей": ["окна кв", "3 окна + балкон", "окна (квартира)"]}),
    "Остекление балкона": (22, (16000, 36000), {
        "Наталья": ["Остекление балкона", "Лоджия под ключ"],
        "Ion": ["balcon", "balcon + izolare"],
        "Андрей": ["балкон", "балкон тёплый", "лоджия"]}),
    "Окно ПВХ": (18, (4500, 9500), {
        "Наталья": ["Окно ПВХ", "Окно двухстворчатое"],
        "Ion": ["fereastră", "fereastră 2 canate"],
        "Андрей": ["окно 2ств", "1 окно кухня", "окно"]}),
    "Входная дверь ПВХ": (10, (9000, 18000), {
        "Наталья": ["Входная дверь ПВХ"],
        "Ion": ["ușă PVC", "ușă intrare"],
        "Андрей": ["дверь", "дверь вход."]}),
    "Частный дом под ключ": (8, (60000, 150000), {
        "Наталья": ["Частный дом, все окна"],
        "Ion": ["casă - toate geamurile", "casă la cheie"],
        "Андрей": ["дом под ключ", "частный дом"]}),
    "Раздвижная система": (6, (35000, 80000), {
        "Наталья": ["Раздвижная система, терраса"],
        "Ion": ["sistem glisant"],
        "Андрей": ["раздвижка", "терраса"]}),
    "Москитные сетки и подоконники": (6, (1200, 4000), {
        "Наталья": ["Москитные сетки", "Подоконники"],
        "Ion": ["plase țânțari", "glafuri"],
        "Андрей": ["сетки", "москитки", "сетки + подоконник"]}),
}
COMMERCIAL = ("Коммерческий объект", (60000, 200000), {
    "Наталья": ["Офис", "Магазин, витрина"],
    "Ion": ["birou", "depozit", "magazin - vitrine"],
    "Андрей": ["офис", "склад", "кафе витрины"]})

# Физлица. Popescu и Ceban исключены: они заняты заложенными ловушками.
RO_FIRST = [("Vasile", "Василе"), ("Maria", "Мария"), ("Elena", "Елена"), ("Victor", "Виктор"),
            ("Mihai", "Михай"), ("Ana", "Ана"), ("Sergiu", "Серджиу"), ("Tatiana", "Татьяна"),
            ("Dumitru", "Думитру"), ("Nicolae", "Николае"), ("Veronica", "Вероника"),
            ("Cristina", "Кристина"), ("Alexandru", "Александру"), ("Doina", "Дойна"),
            ("Radu", "Раду"), ("Liliana", "Лилиана"), ("Valeriu", "Валериу"), ("Aurelia", "Аурелия")]
RO_LAST = [("Rusu", "Русу"), ("Ciobanu", "Чобану"), ("Lungu", "Лунгу"), ("Munteanu", "Мунтяну"),
           ("Rotaru", "Ротару"), ("Cojocaru", "Кожокару"), ("Moraru", "Морару"), ("Botnari", "Ботнарь"),
           ("Țurcanu", "Цуркану"), ("Sârbu", "Сырбу"), ("Guțu", "Гуцу"), ("Bivol", "Бивол"),
           ("Cebotari", "Чеботарь"), ("Ursu", "Урсу"), ("Grosu", "Гросу"), ("Postolachi", "Постолаки"),
           ("Railean", "Райлян"), ("Melnic", "Мельник"), ("Stratan", "Стратан"), ("Lupu", "Лупу"),
           ("Cazacu", "Казаку"), ("Negru", "Негру"), ("Mocanu", "Мокану"), ("Bejan", "Бежан"),
           ("Tabacaru", "Табакару"), ("Oprea", "Опря")]
RU_FIRST_M = [("Сергей", "Serghei"), ("Игорь", "Igor"), ("Дмитрий", "Dmitri"), ("Олег", "Oleg"),
              ("Павел", "Pavel"), ("Виталий", "Vitalie"), ("Юрий", "Iurie"), ("Геннадий", "Ghenadie")]
RU_FIRST_F = [("Светлана", "Svetlana"), ("Ирина", "Irina"), ("Ольга", "Olga"), ("Людмила", "Liudmila"),
              ("Галина", "Galina"), ("Оксана", "Oxana"), ("Алла", "Alla"), ("Лариса", "Larisa")]
RU_LAST = [("Иванов", "Ivanov", True), ("Петренко", "Petrenco", False), ("Ковальчук", "Covalciuc", False),
           ("Кузнецов", "Kuznețov", True), ("Бондаренко", "Bondarenco", False), ("Морозов", "Morozov", True),
           ("Ткаченко", "Tcacenco", False), ("Соколов", "Socolov", True), ("Лысенко", "Lîsenco", False),
           ("Волков", "Volcov", True), ("Шевченко", "Șevcenco", False), ("Козлов", "Cozlov", True),
           ("Попов", "Popov", True), ("Никитин", "Nichitin", True), ("Романов", "Romanov", True),
           ("Захаров", "Zaharov", True), ("Гончаренко", "Goncearenco", False), ("Кравченко", "Cravcenco", False),
           ("Савченко", "Savcenco", False), ("Фёдоров", "Fiodorov", True), ("Белов", "Belov", True)]

# Вымышленные компании (латиница, кириллица). Pomul Prim / Pomul Prim Plus — ловушка.
COMPANIES = [("Stejarel Grup", "Стежарел Груп"), ("Lunca Verde Construct", "Лунка Верде Констракт"),
             ("Colinele Nord", "Колинеле Норд"), ("Zăvoi Prim", "Зэвой Прим"),
             ("Mirabel Studio", "Мирабел Студио"), ("Cireșar Invest", "Чирешар Инвест"),
             ("Brumăriu Service", "Брумэриу Сервис"), ("Tomaș Logistic", "Томаш Логистик")]

REASONS_EARLY = ["передумал", "просто узнать цену", "не наш регион"]
REASONS_MEASURE = ["дорого", "ремонт перенесли", "выбрал конкурентов"]
REASONS_KP = ["дорого", "выбрал конкурентов", "отложили до весны", "передумал", "не берёт трубку"]
COMMENTS = ["звонить после 18", "хочет скидку", "ламинация под дерево", "белые", "жена решает",
            "аванс 50%", "нужен монтаж в субботу", "просил рассрочку", "старые деревянные, демонтаж",
            "4-камерный профиль", "повторный клиент", "sună după ora 18", "vrea reducere", "5 этаж без лифта"]

RU_MON = ["янв", "фев", "мар", "апр", "мая", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"]
RU_MONTH_NAME = {6: "ИЮНЬ", 7: "ИЮЛЬ", 8: "АВГУСТ"}
RU_MONTH_LOWER = {6: "июнь", 7: "июль", 8: "август"}


# ---------------------------------------------------------------------------
# Модель заявки
# ---------------------------------------------------------------------------

@dataclass
class Lead:
    id: int
    manager: str
    channel: str
    kind: str                        # «физлицо» | «компания»
    person: dict | None
    company: tuple | None
    phone: str                       # канон: 9 цифр, начиная с 0
    city: str
    product: str
    lead_date: dt.date
    product_label: str | None = None  # фиксированная подпись товара (для заложенных)
    measure_date: dt.date | None = None
    kp_date: dt.date | None = None
    pay_date: dt.date | None = None
    amount: int | None = None         # MDL
    eur: int | None = None            # если в исходнике сумма в евро
    stage: int = 0
    status: str = ""
    reason: str = ""
    kp_fast: bool | None = None
    planted: str = ""
    issues: list = field(default_factory=list)
    main_ref: str = ""
    dup_refs: list = field(default_factory=list)
    lead_date_known: bool = True
    measure_known: bool = True
    pay_date_known: bool = True
    amount_known: bool = True
    phone_in_source: bool = True

    @property
    def client(self):
        if self.kind == "компания":
            return f"{self.company[0]} SRL"
        return f"{self.person['first_cyr']} {self.person['last_cyr']}"

    @property
    def kp_delay(self):
        if self.measure_date and self.kp_date:
            return (self.kp_date - self.measure_date).days
        return None


# ---------------------------------------------------------------------------
# Генерация
# ---------------------------------------------------------------------------

used_names, used_phones = set(), set()


def pick_weighted(d):
    keys = list(d)
    return rng.choices(keys, weights=[d[k][0] for k in keys])[0]


def make_person():
    while True:
        if rng.random() < 0.6:
            f_lat, f_cyr = rng.choice(RO_FIRST)
            l_lat, l_cyr = rng.choice(RO_LAST)
        else:
            female = rng.random() < 0.5
            f_cyr, f_lat = rng.choice(RU_FIRST_F if female else RU_FIRST_M)
            l_cyr, l_lat, gendered = rng.choice(RU_LAST)
            if female and gendered:
                l_cyr, l_lat = l_cyr + "а", l_lat + "a"
        if (f_cyr, l_cyr) not in used_names:
            used_names.add((f_cyr, l_cyr))
            return dict(first_cyr=f_cyr, last_cyr=l_cyr, first_lat=f_lat, last_lat=l_lat)


def make_phone(landline=False):
    # Диапазоны 060 0xx xxx и 022 0x xx xx выбраны как маловероятные для реальных абонентов.
    while True:
        p = ("0220" if landline else "0600") + f"{rng.randrange(100000):05d}"
        if p not in used_phones:
            used_phones.add(p)
            return p


def random_lead_date():
    span = (PERIOD_END - PERIOD_START).days + 1
    while True:
        d = PERIOD_START + dt.timedelta(days=rng.randrange(span))
        if d.weekday() == 6 and rng.random() < 0.6:
            continue
        return d


def set_amount(lead):
    lo, hi = COMMERCIAL[1] if lead.product == COMMERCIAL[0] else PRODUCTS[lead.product][1]
    n = round(rng.randint(lo, hi), -2)
    eur_share = (0.30 if n >= 30000 else 0.04) if lead.manager == "Ion" else 0.02
    if rng.random() < eur_share:
        lead.eur = max(10, round(n / EUR_RATE / 10) * 10)
        lead.amount = round(lead.eur * EUR_RATE)
    else:
        lead.amount = n


def simulate(lead):
    if rng.random() >= P_MEASURE[lead.channel]:
        lead.stage = 0
        lead.status = "Не дозвонились" if rng.random() < 0.4 else "Отказ"
        if lead.status == "Отказ":
            lead.reason = rng.choice(REASONS_EARLY)
        return
    lead.measure_date = lead.lead_date + dt.timedelta(days=rng.randint(1, 6))
    lead.stage = 1
    if rng.random() >= P_KP_AFTER_MEASURE:
        lead.status, lead.reason = "Отказ", rng.choice(REASONS_MEASURE)
        return
    fast = rng.random() < P_FAST_KP[lead.manager]
    delay = rng.choice([0, 1, 1]) if fast else rng.randint(2, 9)
    lead.kp_date = lead.measure_date + dt.timedelta(days=delay)
    lead.kp_fast = delay <= 1
    lead.stage = 2
    set_amount(lead)
    if rng.random() >= min(0.95, BASE_CONTRACT[lead.channel] * SPEED_FACTOR[fast]):
        recent = (SNAPSHOT - lead.kp_date).days < 10
        lead.status = "Думает" if (recent or rng.random() < 0.35) else "Отказ"
        if lead.status == "Отказ":
            lead.reason = rng.choice(REASONS_KP)
        return
    lead.stage = 3
    pay = lead.kp_date + dt.timedelta(days=rng.randint(2, 16))
    if pay > SNAPSHOT or rng.random() >= P_PAY_AFTER_CONTRACT:
        lead.status = "Договор"
        return
    lead.pay_date, lead.stage, lead.status = pay, 4, "Оплачено"


def generate_leads():
    leads = []
    channels = [c for c, n in CHANNEL_COUNTS.items() for _ in range(n)]
    rng.shuffle(channels)
    companies = COMPANIES[:]
    rng.shuffle(companies)
    managers = list(MANAGERS)
    for ch in channels:
        mgr = rng.choices(managers, weights=[MANAGERS[m] for m in managers])[0]
        is_company = ch in ("Звонок", "Рекомендация") and companies and rng.random() < 0.10
        lead = Lead(
            id=0, manager=mgr, channel=ch,
            kind="компания" if is_company else "физлицо",
            person=None if is_company else make_person(),
            company=companies.pop() if is_company else None,
            phone=make_phone(landline=is_company),
            city=pick_weighted(CITIES),
            product=COMMERCIAL[0] if is_company else pick_weighted(PRODUCTS),
            lead_date=random_lead_date(),
        )
        simulate(lead)
        leads.append(lead)

    # --- заложенные ловушки ---
    pomul = Lead(0, "Андрей", "Звонок", "компания", None, ("Pomul Prim", "Помул Прим"),
                 make_phone(True), "Кишинёв", COMMERCIAL[0], dt.date(2026, 6, 16),
                 product_label="офис, 14 окон", planted="Pomul Prim")
    pomul.measure_date, pomul.kp_date, pomul.pay_date = dt.date(2026, 6, 18), dt.date(2026, 6, 19), dt.date(2026, 7, 2)
    pomul.kp_fast, pomul.amount, pomul.stage, pomul.status = True, 118000, 4, "Оплачено"

    pomul_plus = Lead(0, "Ion", "Рекомендация", "компания", None, ("Pomul Prim Plus", "Помул Прим Плюс"),
                      make_phone(True), "Бельцы", COMMERCIAL[0], dt.date(2026, 7, 7),
                      product_label="depozit: ferestre + uși", planted="Pomul Prim Plus")
    pomul_plus.measure_date, pomul_plus.kp_date, pomul_plus.pay_date = dt.date(2026, 7, 9), dt.date(2026, 7, 15), dt.date(2026, 7, 28)
    pomul_plus.kp_fast, pomul_plus.eur, pomul_plus.stage, pomul_plus.status = False, 4300, 4, "Оплачено"
    pomul_plus.amount = round(4300 * EUR_RATE)

    popescu = Lead(0, "Наталья", "Instagram", "физлицо",
                   dict(first_cyr="Ион", last_cyr="Попеску", first_lat="Ion", last_lat="Popescu"),
                   None, make_phone(), "Кишинёв", "Остекление балкона", dt.date(2026, 7, 13),
                   planted="Ион Попеску")
    popescu.measure_date, popescu.kp_date, popescu.pay_date = dt.date(2026, 7, 15), dt.date(2026, 7, 15), dt.date(2026, 7, 24)
    popescu.kp_fast, popescu.amount, popescu.stage, popescu.status = True, 27400, 4, "Оплачено"

    ceban_1 = Lead(0, "Наталья", "999.md", "физлицо",
                   dict(first_cyr="Мария", last_cyr="Чебан", first_lat="Maria", last_lat="Ceban"),
                   None, make_phone(), "Кишинёв", "Окна в квартиру", dt.date(2026, 6, 24),
                   planted="Мария Чебан (Кишинёв)")
    ceban_1.measure_date, ceban_1.kp_date, ceban_1.pay_date = dt.date(2026, 6, 26), dt.date(2026, 6, 27), dt.date(2026, 7, 8)
    ceban_1.kp_fast, ceban_1.amount, ceban_1.stage, ceban_1.status = True, 31200, 4, "Оплачено"

    ceban_2 = Lead(0, "Андрей", "Рекомендация", "физлицо",
                   dict(first_cyr="Мария", last_cyr="Чебан", first_lat="Maria", last_lat="Ceban"),
                   None, make_phone(), "Орхей", "Входная дверь ПВХ", dt.date(2026, 7, 21),
                   planted="Мария Чебан (Орхей)")
    ceban_2.measure_date, ceban_2.kp_date, ceban_2.pay_date = dt.date(2026, 7, 23), dt.date(2026, 7, 24), dt.date(2026, 8, 5)
    ceban_2.kp_fast, ceban_2.amount, ceban_2.stage, ceban_2.status = True, 14600, 4, "Оплачено"

    leads += [pomul, pomul_plus, popescu, ceban_1, ceban_2]
    leads.sort(key=lambda l: (l.lead_date, l.manager))
    for i, lead in enumerate(leads, 1):
        lead.id = i
    return leads


# ---------------------------------------------------------------------------
# «Загрязнение»: как каждый менеджер записывает одно и то же
# ---------------------------------------------------------------------------

def excel_serial(d):
    return (d - dt.date(1899, 12, 30)).days


def thousands(n, sep=" "):
    return f"{n:,}".replace(",", sep)


def render_phone(p):
    if p.startswith("022"):
        return rng.choice([f"022 {p[3:5]} {p[5:7]} {p[7:]}", f"+373 22 {p[3:]}", f"022-{p[3:6]}-{p[6:]}"])
    return rng.choice([f"{p[:3]} {p[3:6]} {p[6:]}", f"+373{p[1:]}", p[1:],
                       f"{p[:4]}-{p[4:6]}-{p[6:]}", f"+373 {p[1:3]} {p[3:6]} {p[6:]}", p])


def render_date(d, mgr, lead=None, field_name=""):
    if d is None:
        return None
    if mgr == "Наталья":
        r = rng.random()
        if r < 0.85:
            return d
        if r < 0.95:
            return d.strftime("%d.%m.%Y")
        if lead is not None:
            lead.issues.append(f"{field_name}: число вместо даты")
        return excel_serial(d)
    if mgr == "Ion":
        if rng.random() < 0.8:
            return f"{d.month}/{d.day}/{d.year % 100}"   # американский формат: месяц/день/год
        return d
    style = rng.randrange(6)
    return [f"{d.day} {RU_MON[d.month - 1]}", d.strftime("%d.%m"), d.isoformat(),
            d.strftime("%d.%m.%y"), f"{d.day}.{d.month}", d][style]


def render_client(lead, mgr):
    if lead.kind == "компания":
        lat, cyr = lead.company
        if mgr == "Наталья":
            name = f"SRL «{lat}»"
        elif mgr == "Ion":
            name = rng.choice([f"{lat} SRL", f"{lat} S.R.L.", f"SRL {lat}"])
        else:
            name = rng.choice([cyr, f"{lat.lower()} srl"])
    else:
        p = lead.person
        is_ro = any(p["first_lat"] == f for f, _ in RO_FIRST) or p["last_lat"] in ("Popescu", "Ceban")
        if mgr == "Наталья":
            name = f"{p['last_cyr']} {p['first_cyr']}"
        elif mgr == "Ion":
            if is_ro or rng.random() < 0.35:
                name = f"{p['first_lat']} {p['last_lat']}"
            else:
                name = f"{p['first_cyr']} {p['last_cyr']}"
        else:
            name = rng.choice([p["first_cyr"], f"{p['first_cyr']} {p['last_cyr']}", f"{p['first_cyr']} {p['last_cyr'][0]}."])
    return name


def render_amount(lead, mgr):
    n, eur = lead.amount, lead.eur
    if n is None:
        return None
    if lead.planted and eur is None:
        return n
    if eur is not None:
        lead.issues.append("сумма в евро")
        if mgr == "Ion":
            return rng.choice([f"€{thousands(eur)}", f"{eur} eur", f"{thousands(eur, '.')} €", f"{eur} EUR"])
        return f"{eur}€" if mgr == "Андрей" else f"{thousands(eur)} €"
    if mgr == "Наталья":
        r = rng.random()
        if r < 0.80:
            return n
        if r < 0.95:
            return f"{thousands(n, chr(160))} лей"
        return f"{n} MDL"
    if mgr == "Ion":
        return rng.choice([f"{n} lei", thousands(n), n, f"{n},00", f"{thousands(n, '.')} lei"])
    # Андрей
    r = rng.random()
    if r < 0.15 and n > 8000:
        mount = rng.choice([1500, 2000, 2500, 3000])
        lead.issues.append("монтаж вписан в сумму")
        return f"{n - mount} + монтаж {mount}"
    if r < 0.23 and n >= 10000:
        lead.amount = round(n, -3)
        lead.issues.append("сумма примерно («~…к»)")
        return f"~{lead.amount // 1000}к"
    if r < 0.28 and lead.stage < 3:
        lead.amount_known = False
        lead.issues.append("сумма не указана («уточнить»)")
        return "уточнить"
    if r < 0.40 and n >= 1000:
        return thousands(n, ".")
    return n


def render_status(status, mgr):
    return rng.choice(STATUSES[mgr][status if status in STATUSES[mgr] else "Думает"])


def render_product(lead, mgr):
    if lead.product_label:
        return lead.product_label
    variants = COMMERCIAL[2] if lead.product == COMMERCIAL[0] else PRODUCTS[lead.product][2]
    return rng.choice(variants[mgr])


def render_comment(lead, mgr):
    parts = []
    if lead.status == "Отказ" and lead.reason and rng.random() < 0.6:
        parts.append(lead.reason)
    if rng.random() < 0.3:
        parts.append(rng.choice(COMMENTS))
    return ", ".join(parts) or None


def render_row(lead, mgr):
    """Основная строка заявки в стиле менеджера. Возвращает словарь «колонка → значение»."""
    if lead.channel == "Не указан":
        lead.issues.append("источник не указан")
    row = dict(
        date=render_date(lead.lead_date, mgr, lead, "дата заявки"),
        client=render_client(lead, mgr),
        phone=render_phone(lead.phone),
        city=rng.choice(CITIES[lead.city][1][mgr]),
        source=rng.choice(SOURCES[lead.channel][mgr]),
        product=render_product(lead, mgr),
        measure=render_date(lead.measure_date, mgr, lead, "дата замера"),
        kp=render_date(lead.kp_date, mgr, lead, "дата КП"),
        amount=render_amount(lead, mgr),
        status=render_status(lead.status, mgr),
        pay=render_date(lead.pay_date, mgr, lead, "дата оплаты"),
        comment=render_comment(lead, mgr),
    )
    if mgr == "Ion" and any(isinstance(row[k], str) and "/" in row[k] for k in ("date", "measure", "kp", "pay")):
        lead.issues.append("даты в формате месяц/день/год")
    if mgr == "Андрей":
        if lead.measure_date and rng.random() < 0.10:
            row["measure"] = "да"
            lead.measure_known = False
            lead.issues.append("замер: «да» вместо даты")
        if lead.pay_date and rng.random() < 0.35:
            row["pay"] = "+"
            lead.pay_date_known = False
            lead.issues.append("оплата: «+» вместо даты")
        if lead.stage >= 2 and rng.random() < 0.05:
            row["status"] = None
            lead.issues.append("статус пустой")
        if rng.random() < 0.08 and lead.kind == "физлицо" and not lead.planted:
            lead.phone_in_source = False
            lead.issues.append("нет телефона")
            row["client"] = f"{row['client']} ({rng.choice(['Ботаника', 'Чеканы', 'Буюканы', 'центр'])})"
        else:
            row["client"] = f"{row['client']} {row['phone']}"
    return row


def render_shadow(lead, mgr):
    """Повторная запись той же заявки у другого менеджера: клиент позвонил ещё раз."""
    date = lead.lead_date + dt.timedelta(days=rng.randint(0, 3))
    kind = rng.random()
    if kind < 0.4:
        if mgr == "Ion":
            status = f"передал {DATIVE[lead.manager]}"
        elif mgr == "Наталья":
            status = f"Передала {DATIVE[lead.manager]}"
        else:
            status = f"передал {DATIVE[lead.manager]}"
    elif kind < 0.7:
        status = {"Наталья": "Дубль?", "Ion": "dublu?", "Андрей": "дубль?"}[mgr]
    else:
        status = render_status("Не дозвонились", mgr)
    row = dict(
        date=render_date(date, mgr), client=render_client(lead, mgr), phone=render_phone(lead.phone),
        city=rng.choice(CITIES[lead.city][1][mgr]), source=rng.choice(SOURCES[lead.channel][mgr]),
        product=render_product(lead, mgr), measure=None, kp=None, amount=None,
        status=status, pay=None, comment=rng.choice([None, None, "уже общался с нами", "звонил повторно"]),
    )
    if mgr == "Андрей":
        row["client"] = f"{row['client']} {row['phone']}"
    return date, row


# ---------------------------------------------------------------------------
# Сборка грязного файла
# ---------------------------------------------------------------------------

NATALIA_COLS = [("№", None), ("Дата", "date"), ("Клиент", "client"), ("Телефон", "phone"), ("Город", "city"),
                ("Откуда", "source"), ("Что", "product"), ("Замер", "measure"), ("КП", "kp"),
                ("Сумма", "amount"), ("Статус", "status"), ("Оплата", "pay"), ("Комментарий", "comment")]
ION_COLS = [("Nr", None), ("Data", "date"), ("Client", "client"), ("Telefon", "phone"), ("Oraș", "city"),
            ("Sursa", "source"), ("Produs", "product"), ("Suma", "amount"), ("Status", "status"),
            ("Măsurare", "measure"), ("Ofertă", "kp"), ("Achitat", "pay"), ("Note", "comment")]
ANDREI_COLS = [("Дата", "date"), ("Клиент / тел", "client"), ("Город", "city"), ("Источник", "source"),
               ("Заказ", "product"), ("Замер", "measure"), ("КП отпр.", "kp"), ("Сумма", "amount"),
               ("Статус", "status"), ("Оплата", "pay"), ("Примечание", "comment")]

YELLOW = PatternFill("solid", fgColor="FFF2A8")
GREEN = PatternFill("solid", fgColor="C8EFC4")
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")


def write_cell(ws, r, c, value, money=False):
    cell = ws.cell(row=r, column=c, value=value)
    if isinstance(value, dt.date):
        cell.number_format = "DD.MM.YYYY"
    elif money and isinstance(value, int):
        cell.number_format = "#,##0"
    return cell


def build_dirty(leads):
    records = defaultdict(list)   # лист → [(дата для сортировки, строка, заявка, тип)]
    for lead in leads:
        records[lead.manager].append((lead.lead_date, render_row(lead, lead.manager), lead, "main"))

    # Дубли между листами: клиент позвонил ещё раз и попал к другому менеджеру
    candidates = [l for l in leads if not l.planted and l.phone_in_source]
    shadows = rng.sample(candidates, 10)
    shadows += [l for l in leads if l.planted in ("Pomul Prim", "Ион Попеску")]
    for lead in shadows:
        if lead.planted == "Pomul Prim":
            other = "Наталья"
        elif lead.planted == "Ион Попеску":
            other = "Ion"
        else:
            other = rng.choice([m for m in MANAGERS if m != lead.manager])
        date, row = render_shadow(lead, other)
        records[other].append((date, row, lead, "shadow"))
        lead.issues.append(f"дубль в листе «{other}»")

    # Лёгкий беспорядок в порядке строк: менеджеры вписывают заявки задним числом
    for mgr in records:
        recs = sorted(records[mgr], key=lambda x: x[0])
        for _ in range(len(recs) // 20):
            i = rng.randrange(len(recs) - 1)
            recs[i], recs[i + 1] = recs[i + 1], recs[i]
        # Точные копии строк (copy-paste)
        n_copies = {"Наталья": 3, "Ion": 2, "Андрей": 3}[mgr]
        for i in sorted(rng.sample(range(5, len(recs) - 5), n_copies), reverse=True):
            d, row, lead, kind = recs[i]
            if kind == "main":
                recs.insert(i + rng.randint(1, 3), (d, row, lead, "copy"))
                lead.issues.append("точная копия строки")
        records[mgr] = recs

    # «вчера» вместо даты у Андрея
    vchera = [r for r in records["Андрей"] if r[3] == "main" and r[2].stage < 3 and r[2].lead_date.month == 8]
    for d, row, lead, kind in rng.sample(vchera, 2):
        row["date"] = "вчера"
        lead.lead_date_known = False
        lead.issues.append("дата заявки: «вчера»")

    wb = Workbook()
    stats = {}

    # --- Наталья: аккуратнее всех ---
    ws = wb.active
    ws.title = "Наталья"
    for c, (h, _) in enumerate(NATALIA_COLS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill = Font(bold=True), HEADER_FILL
    r, nr, nr_of = 2, 0, {}
    for d, row, lead, kind in records["Наталья"]:
        if kind == "copy":
            num = nr_of[id(row)]
        else:
            nr += 1
            num = nr_of[id(row)] = nr
        write_cell(ws, r, 1, num)
        for c, (_, key) in enumerate(NATALIA_COLS[1:], 2):
            write_cell(ws, r, c, row[key], money=key == "amount")
        register(lead, kind, "Наталья", r)
        r += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{r - 1}"
    set_widths(ws, [5, 12, 26, 16, 22, 15, 26, 12, 12, 14, 16, 12, 30])
    stats["Наталья"] = dict(rows=r - 2, service=0)

    # --- Ion: по-румынски, даты месяц/день/год, суммы в евро ---
    ws = wb.create_sheet("Ion")
    for c, (h, _) in enumerate(ION_COLS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True)
    r, nr, nr_of = 2, 0, {}
    blank_after = len(records["Ion"]) // 2
    for i, (d, row, lead, kind) in enumerate(records["Ion"]):
        if i == blank_after:
            r += 1                     # пустая строка посреди данных
        if kind == "copy":
            num = nr_of[id(row)]
        else:
            nr += 1
            num = nr_of[id(row)] = nr
        write_cell(ws, r, 1, num)
        for c, (_, key) in enumerate(ION_COLS[1:], 2):
            write_cell(ws, r, c, row[key], money=key == "amount")
        register(lead, kind, "Ion", r)
        r += 1
    set_widths(ws, [5, 10, 26, 16, 22, 14, 24, 14, 16, 11, 11, 11, 26])
    stats["Ion"] = dict(rows=len(records["Ion"]), service=1)

    # --- Андрей: заголовок не в первой строке, месяцы, «итого», цвета ---
    ws = wb.create_sheet("Андрей")
    ws.cell(row=1, column=1, value="ПРОДАЖИ АНДРЕЙ — ЛЕТО 2026").font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(ANDREI_COLS))
    ws.cell(row=1, column=13, value="жёлтый = перезвонить, зелёный = оплачено (не всегда)").font = Font(italic=True)
    for c, (h, _) in enumerate(ANDREI_COLS, 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font = Font(bold=True)
    r, service = 4, 0
    andrei_subtotals = {}
    by_month = defaultdict(list)
    for rec in records["Андрей"]:
        by_month[rec[0].month].append(rec)
    for month in sorted(by_month):
        cell = ws.cell(row=r, column=1, value=RU_MONTH_NAME[month])
        cell.font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(ANDREI_COLS))
        cell.alignment = Alignment(horizontal="center")
        r += 1
        service += 1
        month_sum = 0
        for d, row, lead, kind in by_month[month]:
            for c, (_, key) in enumerate(ANDREI_COLS, 1):
                write_cell(ws, r, c, row[key], money=key == "amount")
            fill = None
            if lead.status == "Оплачено" and kind != "shadow" and rng.random() < 0.6:
                fill = GREEN
            elif lead.status == "Думает" and kind != "shadow" and rng.random() < 0.5:
                fill = YELLOW
            if fill:
                for c in range(1, len(ANDREI_COLS) + 1):
                    ws.cell(row=r, column=c).fill = fill
            if kind == "main" and lead.amount and lead.amount_known:
                month_sum += lead.amount    # Андрей суммирует все КП, а не оплаты
            register(lead, kind, "Андрей", r)
            r += 1
        ws.cell(row=r, column=2, value=f"Итого {RU_MONTH_LOWER[month]}").font = Font(bold=True)
        write_cell(ws, r, 8, month_sum, money=True).font = Font(bold=True)
        andrei_subtotals[month] = (r, month_sum)
        r += 2
        service += 1
    set_widths(ws, [11, 30, 16, 16, 18, 10, 10, 18, 16, 10, 26])
    stats["Андрей"] = dict(rows=len(records["Андрей"]), service=service + 2, subtotals=andrei_subtotals)

    # --- Свод бухгалтера: начатый и брошенный ---
    ws = wb.create_sheet("Свод Марина")
    ws["A1"] = "СВОД ИЮНЬ"
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Менеджер", "Оплат", "Сумма"])
    june_paid = defaultdict(lambda: [0, 0])
    for lead in leads:
        if lead.status == "Оплачено" and lead.pay_date and lead.pay_date.month == 6:
            june_paid[lead.manager][0] += 1
            june_paid[lead.manager][1] += lead.amount
    for mgr, label in (("Наталья", "Наталья"), ("Ion", "Ион"), ("Андрей", "Андрей")):
        cnt, total = june_paid[mgr]
        if mgr == "Андрей":
            total = swap_digits(total)   # ручная ошибка: переставлены цифры
        ws.append([label, cnt, total])
    ws.append(["ИТОГО", "=SUM(B4:B6)", "=SUM(C4:C6)"])
    ws.append([])
    ws.append(["июль, август — не успела, доделать!!!"])
    ws["A9"].font = Font(color="C00000", bold=True)
    set_widths(ws, [40, 10, 14])

    wb.save(DIRTY_FILE)
    return stats


def swap_digits(n):
    """57100 → 51700: типичная ошибка ручного ввода."""
    d = list(str(n))
    for i in range(1, len(d) - 1):
        if d[i] != d[i + 1]:
            d[i], d[i + 1] = d[i + 1], d[i]
            break
    return int("".join(d))


def register(lead, kind, sheet, row):
    ref = f"{sheet}!{row}"
    if kind == "main":
        lead.main_ref = ref
    else:
        lead.dup_refs.append(ref + (" (копия)" if kind == "copy" else " (повторная заявка)"))


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------------------------------------------------------------------
# Эталон
# ---------------------------------------------------------------------------

def plural(n, one, few, many):
    if n % 10 == 1 and n % 100 != 11:
        return f"{n} {one}"
    if 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14:
        return f"{n} {few}"
    return f"{n} {many}"


def pct(a, b):
    return round(100 * a / b, 1) if b else 0.0


def funnel(group):
    counts = [sum(1 for l in group if l.stage >= s) for s in range(5)]
    revenue = sum(l.amount for l in group if l.stage == 4)
    return counts, revenue


def build_truth(leads, stats):
    wb = Workbook()
    bold = Font(bold=True)

    # --- Заявки ---
    ws = wb.active
    ws.title = "Заявки"
    head = ["ID", "Менеджер", "Дата заявки", "Клиент", "Тип", "Телефон", "Город", "Источник", "Товар",
            "Сумма, MDL", "Было в EUR", "Дата замера", "Дата КП", "Дней замер→КП", "КП за сутки",
            "Дата оплаты", "Этап", "Статус", "Причина отказа", "Основная строка", "Дубли",
            "Особенности исходника", "Ловушка"]
    ws.append(head)
    for c in range(1, len(head) + 1):
        ws.cell(row=1, column=c).font = bold
    for l in leads:
        ws.append([l.id, l.manager, l.lead_date, l.client, l.kind, l.phone, l.city, l.channel,
                   l.product, l.amount if l.amount_known else None, l.eur,
                   l.measure_date if l.measure_known else None, l.kp_date,
                   l.kp_delay if l.measure_known else None,
                   ("да" if l.kp_fast else "нет") if (l.kp_fast is not None and l.measure_known) else None,
                   l.pay_date if l.pay_date_known else None, STAGES[l.stage], l.status, l.reason or None,
                   l.main_ref, "; ".join(l.dup_refs) or None, "; ".join(dict.fromkeys(l.issues)) or None,
                   l.planted or None])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, dt.date):
                cell.number_format = "DD.MM.YYYY"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    set_widths(ws, [5, 10, 12, 24, 10, 12, 12, 13, 26, 12, 9, 12, 12, 9, 9, 12, 9, 14, 18, 14, 36, 50, 22])

    revenue = sum(l.amount for l in leads if l.stage == 4)
    eur_revenue = sum(l.amount for l in leads if l.stage == 4 and l.eur)

    # --- Сверка ---
    ws = wb.create_sheet("Сверка")
    total_rows = sum(s["rows"] for s in stats.values())
    copies = sum(1 for l in leads for d in l.dup_refs if "копия" in d)
    shadows = sum(1 for l in leads for d in l.dup_refs if "повторная" in d)
    rows = [
        ["Показатель", "Значение", "Комментарий"],
        ["Строк с заявками в листах менеджеров", total_rows, "без заголовков, строк месяцев, «Итого» и пустых"],
        ["  лист «Наталья»", stats["Наталья"]["rows"], ""],
        ["  лист «Ion»", stats["Ion"]["rows"], "плюс 1 пустая строка посреди данных"],
        ["  лист «Андрей»", stats["Андрей"]["rows"], f"плюс {stats['Андрей']['service']} служебных строк: заголовок, месяцы, «Итого»"],
        ["Точные копии строк (copy-paste)", copies, "удалить"],
        ["Повторные заявки того же клиента у другого менеджера", shadows, "склеить с основной записью по телефону"],
        ["Уникальных заявок", len(leads), "= строк − копии − повторные"],
        ["Оплачено сделок", sum(1 for l in leads if l.stage == 4), ""],
        ["Выручка (оплачено), MDL", revenue, f"контрольная цифра «поступления по банку» с 01.06 по {SNAPSHOT:%d.%m.%Y}"],
        ["  в т.ч. из сумм в EUR, MDL", eur_revenue, f"по курсу {EUR_RATE}"],
        ["Лист «Свод Марина»", "не данные", "начатый вручную свод за июнь, игнорировать или сверить"],
    ]
    for m, (row_no, s) in stats["Андрей"]["subtotals"].items():
        true_paid = sum(l.amount for l in leads if l.manager == "Андрей" and l.stage == 4 and l.lead_date.month == m)
        rows.append([f"«Итого {RU_MONTH_LOWER[m]}» у Андрея (строка {row_no})", s,
                     f"это сумма всех КП; оплачено по заявкам месяца на самом деле {thousands(true_paid)}"])
    for r in rows:
        ws.append(r)
    ws["A1"].font = ws["B1"].font = ws["C1"].font = bold
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            if isinstance(cell.value, int):
                cell.number_format = "#,##0"
    set_widths(ws, [52, 14, 80])

    # --- Воронки ---
    def funnel_sheet(title, key, order):
        ws = wb.create_sheet(title)
        head = [key, *STAGES, "Заявка→оплата, %", "Замер→КП, %", "КП→договор, %", "Выручка, MDL", "Средний чек, MDL"]
        ws.append(head)
        for c in range(1, len(head) + 1):
            ws.cell(row=1, column=c).font = bold
        groups = defaultdict(list)
        for l in leads:
            groups[l.channel if key == "Канал" else l.manager].append(l)
        for g in [*order, "Итого"]:
            group = leads if g == "Итого" else groups[g]
            (a, b, c, d, e), rev = funnel(group)
            ws.append([g, a, b, c, d, e, pct(e, a), pct(c, b), pct(d, c), rev, round(rev / e) if e else 0])
        for row in ws.iter_rows(min_row=2, min_col=10, max_col=11):
            for cell in row:
                cell.number_format = "#,##0"
        set_widths(ws, [16, 9, 9, 9, 9, 9, 16, 13, 14, 14, 16])
        return ws

    funnel_sheet("Воронка по каналам", "Канал", list(CHANNEL_COUNTS))
    ws = funnel_sheet("Воронка по менеджерам", "Менеджер", list(MANAGERS))
    ws.cell(row=1, column=12, value="КП за сутки, % от КП").font = bold
    for i, m in enumerate(MANAGERS, 2):
        kps = [l for l in leads if l.manager == m and l.kp_fast is not None and l.measure_known]
        ws.cell(row=i, column=12, value=pct(sum(1 for l in kps if l.kp_fast), len(kps)))

    # --- Скорость КП ---
    ws = wb.create_sheet("Скорость КП")
    kps = [l for l in leads if l.kp_date and l.measure_known]
    fast = [l for l in kps if l.kp_fast]
    slow = [l for l in kps if not l.kp_fast]
    conv_f = sum(1 for l in fast if l.stage >= 3) / len(fast)
    conv_s = sum(1 for l in slow if l.stage >= 3) / len(slow)
    contracts = [l for l in leads if l.stage >= 3 and l.amount_known]
    avg_contract = sum(l.amount for l in contracts) / len(contracts)
    extra = len(slow) * (conv_f - conv_s)
    lost_money = extra * avg_contract
    ws.append(["", "КП отправлено", "Дошли до договора", "Конверсия КП→договор, %"])
    ws.append(["КП в течение суток после замера", len(fast), sum(1 for l in fast if l.stage >= 3), round(100 * conv_f, 1)])
    ws.append(["КП позже", len(slow), sum(1 for l in slow if l.stage >= 3), round(100 * conv_s, 1)])
    ws.append([])
    ws.append(["Во сколько раз выше конверсия быстрых КП", round(conv_f / conv_s, 2)])
    ws.append(["Средний чек договора, MDL", round(avg_contract)])
    ws.append(["Если бы медленные КП ушли за сутки: доп. договоров за 3 месяца", round(extra, 1)])
    ws.append(["Упущенная выручка за 3 месяца, MDL", int(round(lost_money, -3))])
    ws.append(["Упущенная выручка в месяц, MDL", int(round(lost_money / MONTHS_IN_PERIOD, -3))])
    ws.append([])
    ws.append(["Внутри каждого менеджера", "Быстрые КП: конверсия, %", "Медленные КП: конверсия, %", "Доля быстрых КП, %"])
    for m in MANAGERS:
        f_m = [l for l in fast if l.manager == m]
        s_m = [l for l in slow if l.manager == m]
        ws.append([m, pct(sum(1 for l in f_m if l.stage >= 3), len(f_m)), pct(sum(1 for l in s_m if l.stage >= 3), len(s_m)),
                   pct(len(f_m), len(f_m) + len(s_m))])
    ws.append([])
    ws.append(["Без учёта заявок, где у Андрея в замере «да» вместо даты. Оценка упущенной выручки — верхняя граница: "
               "считаем, что медленные КП конвертировались бы как быстрые."])
    for c in range(1, 5):
        ws.cell(row=1, column=c).font = bold
    set_widths(ws, [62, 16, 18, 24])

    # --- Ловушки ---
    ws = wb.create_sheet("Ловушки")
    ws.append(["Ловушка", "Где", "Правильно", "Если ошибиться"])
    for c in range(1, 5):
        ws.cell(row=1, column=c).font = bold
    by_plant = {l.planted: l for l in leads if l.planted}
    pp, ppp = by_plant["Pomul Prim"], by_plant["Pomul Prim Plus"]
    ws.append(["Pomul Prim и Pomul Prim Plus", f"{pp.main_ref}, {ppp.main_ref}; повтор Pomul Prim: {', '.join(pp.dup_refs)}",
               "Это две разные компании: разные города и телефоны. Pomul Prim из листа «Наталья» — повтор заявки Андрея.",
               f"Склейка теряет {thousands(ppp.amount)} MDL выручки, итог не сойдётся с банком"])
    pop = by_plant["Ион Попеску"]
    ws.append(["Ион Попеску / Ion Popescu", f"{pop.main_ref}; {', '.join(pop.dup_refs)}",
               "Один человек: тот же телефон, кириллица и латиница", "Двойной счёт заявки"])
    c1, c2 = by_plant["Мария Чебан (Кишинёв)"], by_plant["Мария Чебан (Орхей)"]
    ws.append(["Две Марии Чебан", f"{c1.main_ref}, {c2.main_ref}",
               "Разные люди: разные города и телефоны", f"Склейка теряет {thousands(c2.amount)} MDL выручки"])
    ws.append(["Суммы в EUR", "в основном лист «Ion»", f"Перевести по курсу {EUR_RATE}",
               f"Если считать евро как леи, выручка занижена примерно на {thousands(round(eur_revenue - eur_revenue / EUR_RATE))} MDL"])
    ws.append(["Даты у Иона: месяц/день/год", "лист «Ion»", "7/3/26 — это 3 июля. Март вне периода, так что формат однозначен",
               "Заявки «уедут» в март–декабрь, воронка по месяцам сломается"])
    ws.append(["«Итого июнь/июль/август» у Андрея", "лист «Андрей»", "Служебные строки, исключить",
               "Выручка раздуется на сумму всех КП"])
    june_andrei = sum(l.amount for l in leads if l.manager == "Андрей" and l.stage == 4 and l.pay_date.month == 6)
    ws.append(["«Свод Марина»", "отдельный лист",
               f"Не данные: ручной свод за июнь. У Андрея {thousands(swap_digits(june_andrei))} вместо {thousands(june_andrei)} (переставлены цифры)",
               "Двойной счёт или неверная сверка"])
    ws.append(["«вчера» вместо даты", "лист «Андрей», 2 строки", "Пометить для человека, не выдумывать дату",
               "Выдуманная дата"])
    ws.append(["«уточнить», «~20к», «18500 + монтаж 2000»", "лист «Андрей»",
               "уточнить → пусто и пометить; ~20к → 20 000 и пометить; «+ монтаж» → сложить", "Потеря или искажение сумм"])
    ws.append(["Точные копии и повторные заявки", "все листы", "Копии удалить, повторные склеить по телефону",
               "Раздутое число заявок, заниженная конверсия"])
    set_widths(ws, [36, 44, 70, 56])

    # --- Для слайдов ---
    ws = wb.create_sheet("Для слайдов")
    ch = defaultdict(list)
    for l in leads:
        ch[l.channel].append(l)
    (ia, *_, ie), _ = funnel(ch["Instagram"])
    (ra, *_, re_), _ = funnel(ch["Рекомендация"])
    (_, b, c, d, _), _ = funnel(leads)
    mg = {m: funnel([l for l in leads if l.manager == m]) for m in MANAGERS}
    lines = [
        f"Из Instagram {plural(ia, 'заявка', 'заявки', 'заявок')} и {plural(ie, 'оплата', 'оплаты', 'оплат')} ({pct(ie, ia)}%), "
        f"из рекомендаций {plural(ra, 'заявка', 'заявки', 'заявок')} и {plural(re_, 'оплата', 'оплаты', 'оплат')} ({pct(re_, ra)}%).",
        f"КП отправлено {c}, до договора дошли {d}: на этапе КП → договор теряется {round(100 - pct(d, c))}%.",
        f"КП в течение суток после замера: конверсия в договор {round(100 * conv_f)}%, позже: {round(100 * conv_s)}%. Разница в {conv_f / conv_s:.1f} раза.",
        f"Медленные КП стоят примерно {thousands(int(round(lost_money / MONTHS_IN_PERIOD, -3)))} лей в месяц.",
        "Менеджеры, заявка → оплата: " + ", ".join(f"{m} {pct(v[0][4], v[0][0])}%" for m, v in mg.items()) + ".",
        f"Строк в листах менеджеров: {total_rows}. Уникальных заявок: {len(leads)}. Убрано копий: {copies}, склеено повторных: {shadows}.",
        f"Выручка за лето (оплачено): {thousands(revenue)} MDL.",
    ]
    for line in lines:
        ws.append([line])
    set_widths(ws, [140])

    wb.save(TRUTH_FILE)
    return lines


def main():
    leads = generate_leads()
    stats = build_dirty(leads)
    lines = build_truth(leads, stats)
    print(f"Грязный файл: {DIRTY_FILE.name}")
    print(f"Эталон:       {TRUTH_FILE.name}")
    print()
    for line in lines:
        print("•", line)


if __name__ == "__main__":
    main()
